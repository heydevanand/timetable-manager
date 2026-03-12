import os
import sys
import io
import logging
import tempfile
from datetime import datetime
from functools import wraps
from flask import Flask, render_template, redirect, url_for, flash, request, session, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from flask_wtf.csrf import CSRFProtect, generate_csrf
from flask_session import Session

# Setup logging
logging.basicConfig(
    stream=sys.stdout,
    level=logging.DEBUG,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)

# Initialize Flask app
app = Flask(__name__)

# Configure app
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///timetable.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SESSION_TYPE'] = 'filesystem'
app.config['SESSION_FILE_DIR'] = os.path.join(tempfile.gettempdir(), 'flask_session')
app.config['SESSION_PERMANENT'] = True
app.config['SESSION_USE_SIGNER'] = True
app.config['PERMANENT_SESSION_LIFETIME'] = 3600  # 1 hour
app.config['WTF_CSRF_ENABLED'] = True
app.config['WTF_CSRF_TIME_LIMIT'] = None  # No expiration for CSRF tokens

# Load or generate secret key
SECRET_KEY_FILE = os.path.join(app.root_path, 'secret_key')
if os.path.exists(SECRET_KEY_FILE):
    with open(SECRET_KEY_FILE, 'rb') as f:
        app.config['SECRET_KEY'] = f.read()
else:
    app.config['SECRET_KEY'] = os.urandom(24)
    with open(SECRET_KEY_FILE, 'wb') as f:
        f.write(app.config['SECRET_KEY'])

# Create session directory if it doesn't exist
os.makedirs(app.config['SESSION_FILE_DIR'], exist_ok=True)

# Initialize extensions
db = SQLAlchemy(app)
csrf = CSRFProtect(app)
Session(app)
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'
login_manager.login_message_category = 'info'

# ─── Constants ───────────────────────────────────────────────────────────────

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
TIME_SLOTS = [
    "08:00 - 09:00", "09:00 - 10:00", "10:00 - 11:00", "11:00 - 12:00",
    "12:00 - 01:00", "01:00 - 02:00", "02:00 - 03:00", "03:00 - 04:00",
    "04:00 - 05:00", "05:00 - 06:00"
]

# ─── Models ──────────────────────────────────────────────────────────────────

class User(db.Model, UserMixin):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(50), unique=True, nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(256), nullable=False)
    is_admin = db.Column(db.Boolean, default=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    full_name = db.Column(db.String(100), nullable=True)
    department = db.Column(db.String(100), nullable=True)

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)

    def __repr__(self):
        return f"User('{self.username}', '{self.email}')"

class Teacher(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False, unique=True)
    department = db.Column(db.String(100), nullable=True)
    email = db.Column(db.String(120), nullable=True)
    phone = db.Column(db.String(20), nullable=True)
    schedules = db.relationship('Schedule', backref='teacher', lazy=True, cascade="all, delete-orphan")

    def __repr__(self):
        return f"Teacher('{self.name}')"

class Schedule(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    day = db.Column(db.String(20), nullable=False)
    time_slot = db.Column(db.String(20), nullable=False)
    status = db.Column(db.String(10), nullable=False, default='Free')
    room_number = db.Column(db.String(20), nullable=True)
    subject = db.Column(db.String(100), nullable=True)
    teacher_id = db.Column(db.Integer, db.ForeignKey('teacher.id'), nullable=False)

    def __repr__(self):
        return f"Schedule('{self.day}', '{self.time_slot}', '{self.status}', '{self.room_number}')"

# ─── User Loader ─────────────────────────────────────────────────────────────

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

# ─── Utility Functions ───────────────────────────────────────────────────────

def create_admin_user():
    """Create or reset the admin user."""
    admin = User.query.filter_by(username='admin').first()
    if admin:
        admin.set_password('admin123')
    else:
        admin = User(
            username='admin',
            email='admin@example.com',
            is_admin=True,
            full_name='Administrator',
            department='IT'
        )
        admin.set_password('admin123')
        db.session.add(admin)

    db.session.commit()
    return admin

def init_data():
    """Initialize database with sample data if empty."""
    if Teacher.query.first():
        return

    teachers_data = [
        ("Ashish Mishra", "Computer Science", "ashish@example.com"),
        ("Himanshu", "Mathematics", "himanshu@example.com"),
        ("Dev Anand", "Computer Science", "dev@example.com"),
        ("Harish Ojha", "Physics", "harish@example.com"),
        ("Yogesh Vajpayee", "Mathematics", "yogesh@example.com"),
        ("Shubham", "Electronics", "shubham@example.com"),
        ("Mukesh Jangid", "Physics", "mukesh@example.com"),
        ("Kapil Manchandani", "Electronics", "kapil@example.com"),
    ]

    timetable_data = [
        ("Ashish Mishra", "Monday", "08:00 - 09:00", "Engaged", "201", "Data Structures"),
        ("Himanshu", "Monday", "08:00 - 09:00", "Engaged", "107", "Calculus I"),
        ("Dev Anand", "Monday", "09:00 - 10:00", "Engaged", "201", "Python Programming"),
        ("Harish Ojha", "Monday", "10:00 - 11:00", "Engaged", "305", "Quantum Physics"),
        ("Ashish Mishra", "Tuesday", "09:00 - 10:00", "Engaged", "201", "Algorithms"),
        ("Yogesh Vajpayee", "Tuesday", "08:00 - 09:00", "Engaged", "102", "Linear Algebra"),
        ("Shubham", "Wednesday", "10:00 - 11:00", "Engaged", "Lab-1", "Digital Electronics"),
        ("Mukesh Jangid", "Wednesday", "08:00 - 09:00", "Engaged", "305", "Mechanics"),
        ("Kapil Manchandani", "Thursday", "11:00 - 12:00", "Engaged", "Lab-2", "Circuit Design"),
        ("Himanshu", "Thursday", "09:00 - 10:00", "Engaged", "102", "Probability"),
        ("Dev Anand", "Friday", "08:00 - 09:00", "Engaged", "201", "Database Systems"),
        ("Harish Ojha", "Friday", "10:00 - 11:00", "Engaged", "305", "Thermodynamics"),
    ]

    teacher_objects = {}
    for name, dept, email in teachers_data:
        teacher = Teacher(name=name, department=dept, email=email)
        db.session.add(teacher)
        db.session.commit()
        teacher_objects[name] = teacher

    for teacher_name in teacher_objects:
        teacher = teacher_objects[teacher_name]
        for day in DAYS:
            for time_slot in TIME_SLOTS:
                schedule = Schedule(
                    day=day,
                    time_slot=time_slot,
                    status="Free",
                    room_number=None,
                    subject=None,
                    teacher=teacher
                )
                db.session.add(schedule)

    for data in timetable_data:
        teacher_name, day, time_slot, status, room_number, subject = data
        teacher = teacher_objects[teacher_name]
        schedule = Schedule.query.filter_by(
            teacher_id=teacher.id,
            day=day,
            time_slot=time_slot
        ).first()
        if schedule:
            schedule.status = status
            schedule.room_number = room_number
            schedule.subject = subject

    db.session.commit()
    create_admin_user()

# ─── Decorators ──────────────────────────────────────────────────────────────

def admin_required(f):
    @wraps(f)
    @login_required
    def decorated_function(*args, **kwargs):
        if not current_user.is_admin:
            flash('You do not have permission to access this page.', 'danger')
            return redirect(url_for('index'))
        return f(*args, **kwargs)
    return decorated_function

# ─── Context Processors ─────────────────────────────────────────────────────

@app.context_processor
def inject_csrf_token():
    return dict(csrf_token=generate_csrf())

@app.context_processor
def inject_constants():
    return dict(DAYS=DAYS, TIME_SLOTS=TIME_SLOTS)

# ─── Auth Routes ─────────────────────────────────────────────────────────────

@app.route('/login', methods=['GET', 'POST'])
@csrf.exempt
def login():
    if current_user.is_authenticated:
        return redirect(url_for('index'))

    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        user = User.query.filter_by(username=username).first()

        if not user:
            flash('Login failed. User does not exist.', 'danger')
        elif not user.check_password(password):
            flash('Login failed. Incorrect password.', 'danger')
        else:
            session.permanent = True
            login_user(user, remember=('remember' in request.form))
            next_page = request.args.get('next')
            flash('Login successful!', 'success')
            return redirect(next_page or url_for('index'))

    return render_template('login.html')

@app.route('/register', methods=['GET', 'POST'])
@csrf.exempt
def register():
    if current_user.is_authenticated:
        return redirect(url_for('index'))

    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        email = request.form.get('email', '').strip()
        password = request.form.get('password', '')
        confirm_password = request.form.get('confirm_password', '')

        if User.query.filter_by(username=username).first():
            flash('Username already exists.', 'danger')
            return redirect(url_for('register'))

        if User.query.filter_by(email=email).first():
            flash('Email already registered.', 'danger')
            return redirect(url_for('register'))

        if password != confirm_password:
            flash('Passwords do not match.', 'danger')
            return redirect(url_for('register'))

        if len(password) < 6:
            flash('Password must be at least 6 characters.', 'danger')
            return redirect(url_for('register'))

        user = User(username=username, email=email)
        user.set_password(password)
        db.session.add(user)
        db.session.commit()

        flash('Registration successful! You can now log in.', 'success')
        return redirect(url_for('login'))

    return render_template('register.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('You have been logged out.', 'info')
    return redirect(url_for('login'))

@app.route('/profile')
@login_required
def profile():
    return render_template('profile.html')

@app.route('/profile/edit', methods=['GET', 'POST'])
@login_required
def edit_profile():
    if request.method == 'POST':
        email = request.form.get('email', '').strip()
        full_name = request.form.get('full_name', '').strip()
        department = request.form.get('department', '').strip()
        current_password = request.form.get('current_password', '')
        new_password = request.form.get('new_password', '')
        confirm_password = request.form.get('confirm_password', '')

        if not current_user.check_password(current_password):
            flash('Current password is incorrect', 'danger')
            return redirect(url_for('edit_profile'))

        if email != current_user.email:
            existing = User.query.filter_by(email=email).first()
            if existing and existing.id != current_user.id:
                flash('Email is already in use', 'danger')
                return redirect(url_for('edit_profile'))

        current_user.email = email
        current_user.full_name = full_name
        current_user.department = department

        if new_password:
            if new_password != confirm_password:
                flash('New passwords do not match', 'danger')
                return redirect(url_for('edit_profile'))
            if len(new_password) < 6:
                flash('New password must be at least 6 characters', 'danger')
                return redirect(url_for('edit_profile'))
            current_user.set_password(new_password)

        db.session.commit()
        flash('Profile updated successfully', 'success')
        return redirect(url_for('profile'))

    return render_template('edit_profile.html')

# ─── Main Routes ─────────────────────────────────────────────────────────────

@app.route('/')
def index():
    if not current_user.is_authenticated:
        return redirect(url_for('login'))
    teachers = Teacher.query.order_by(Teacher.name).all()
    departments = db.session.query(Teacher.department).filter(
        Teacher.department.isnot(None)
    ).distinct().order_by(Teacher.department).all()
    departments = [d[0] for d in departments]
    return render_template('index.html', teachers=teachers, departments=departments)

@app.route('/dashboard')
@login_required
def dashboard():
    teachers = Teacher.query.all()
    total_teachers = len(teachers)
    total_schedules = Schedule.query.count()
    engaged_count = Schedule.query.filter_by(status='Engaged').count()
    free_count = Schedule.query.filter_by(status='Free').count()

    # Utilization percentage
    utilization = round((engaged_count / total_schedules * 100), 1) if total_schedules > 0 else 0

    # Per-teacher utilization
    teacher_stats = []
    for teacher in teachers:
        total = len(teacher.schedules)
        engaged = sum(1 for s in teacher.schedules if s.status == 'Engaged')
        pct = round((engaged / total * 100), 1) if total > 0 else 0
        teacher_stats.append({
            'id': teacher.id,
            'name': teacher.name,
            'department': teacher.department or 'N/A',
            'total': total,
            'engaged': engaged,
            'free': total - engaged,
            'utilization': pct
        })
    teacher_stats.sort(key=lambda x: x['utilization'], reverse=True)

    # Room usage stats
    rooms = db.session.query(
        Schedule.room_number,
        db.func.count(Schedule.id)
    ).filter(
        Schedule.room_number.isnot(None),
        Schedule.status == 'Engaged'
    ).group_by(Schedule.room_number).order_by(
        db.func.count(Schedule.id).desc()
    ).all()

    # Day-wise engagement
    day_stats = []
    for day in DAYS:
        engaged = Schedule.query.filter_by(day=day, status='Engaged').count()
        total = Schedule.query.filter_by(day=day).count()
        pct = round((engaged / total * 100), 1) if total > 0 else 0
        day_stats.append({'day': day, 'engaged': engaged, 'total': total, 'utilization': pct})

    # Department stats
    dept_stats = db.session.query(
        Teacher.department,
        db.func.count(db.distinct(Teacher.id))
    ).filter(
        Teacher.department.isnot(None)
    ).group_by(Teacher.department).all()

    return render_template('dashboard.html',
                           total_teachers=total_teachers,
                           total_schedules=total_schedules,
                           engaged_count=engaged_count,
                           free_count=free_count,
                           utilization=utilization,
                           teacher_stats=teacher_stats,
                           rooms=rooms,
                           day_stats=day_stats,
                           dept_stats=dept_stats)

@app.route('/timetable')
@login_required
def timetable():
    teachers = Teacher.query.order_by(Teacher.name).all()
    return render_template('timetable.html', days=DAYS, time_slots=TIME_SLOTS, teachers=teachers)

@app.route('/teacher/<int:id>')
@login_required
def teacher_timetable(id):
    teacher = Teacher.query.get_or_404(id)

    # Build schedule grid
    schedule_grid = {}
    for day in DAYS:
        schedule_grid[day] = {}
        for ts in TIME_SLOTS:
            schedule = Schedule.query.filter_by(
                teacher_id=teacher.id, day=day, time_slot=ts
            ).first()
            schedule_grid[day][ts] = schedule

    # Calculate stats
    total = len(teacher.schedules)
    engaged = sum(1 for s in teacher.schedules if s.status == 'Engaged')
    free = total - engaged
    utilization = round((engaged / total * 100), 1) if total > 0 else 0

    # Subjects taught
    subjects = db.session.query(Schedule.subject).filter(
        Schedule.teacher_id == teacher.id,
        Schedule.subject.isnot(None)
    ).distinct().all()
    subjects = [s[0] for s in subjects]

    # Rooms used
    rooms_used = db.session.query(Schedule.room_number).filter(
        Schedule.teacher_id == teacher.id,
        Schedule.room_number.isnot(None)
    ).distinct().all()
    rooms_used = [r[0] for r in rooms_used]

    return render_template('teacher_timetable.html',
                           teacher=teacher,
                           schedule_grid=schedule_grid,
                           days=DAYS,
                           time_slots=TIME_SLOTS,
                           total=total,
                           engaged=engaged,
                           free=free,
                           utilization=utilization,
                           subjects=subjects,
                           rooms_used=rooms_used)

@app.route('/free_teachers')
@login_required
def free_teachers():
    selected_day = request.args.get('day', DAYS[0])
    selected_time = request.args.get('time_slot', TIME_SLOTS[0])

    free_teachers_list = Teacher.query.join(Schedule).filter(
        Schedule.day == selected_day,
        Schedule.time_slot == selected_time,
        Schedule.status == 'Free'
    ).order_by(Teacher.name).all()

    engaged_teachers_list = Teacher.query.join(Schedule).filter(
        Schedule.day == selected_day,
        Schedule.time_slot == selected_time,
        Schedule.status == 'Engaged'
    ).order_by(Teacher.name).all()

    return render_template('free_teachers.html',
                           free_teachers=free_teachers_list,
                           engaged_teachers=engaged_teachers_list,
                           days=DAYS,
                           time_slots=TIME_SLOTS,
                           selected_day=selected_day,
                           selected_time=selected_time)

# ─── Export Routes ───────────────────────────────────────────────────────────

@app.route('/export/timetable')
@login_required
def export_timetable():
    """Export the full timetable as an Excel file."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    teachers = Teacher.query.order_by(Teacher.name).all()

    for day_idx, day in enumerate(DAYS):
        if day_idx == 0:
            ws = wb.active
            ws.title = day
        else:
            ws = wb.create_sheet(title=day)

        # Header row
        ws.cell(row=1, column=1, value="Time Slot").font = Font(bold=True, size=11)
        ws.cell(row=1, column=1).fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        ws.cell(row=1, column=1).font = Font(bold=True, color="FFFFFF", size=11)
        ws.cell(row=1, column=1).alignment = Alignment(horizontal='center')
        ws.cell(row=1, column=1).border = thin_border
        ws.column_dimensions['A'].width = 16

        for col, teacher in enumerate(teachers, start=2):
            cell = ws.cell(row=1, column=col, value=teacher.name)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
            cell.border = thin_border
            ws.column_dimensions[cell.column_letter].width = 18

        # Data rows
        for row_idx, ts in enumerate(TIME_SLOTS, start=2):
            ws.cell(row=row_idx, column=1, value=ts).font = Font(bold=True, size=10)
            ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=1).border = thin_border

            for col, teacher in enumerate(teachers, start=2):
                schedule = Schedule.query.filter_by(
                    teacher_id=teacher.id, day=day, time_slot=ts
                ).first()

                if schedule and schedule.status == 'Engaged':
                    text = schedule.subject or 'Engaged'
                    if schedule.room_number:
                        text += f"\n({schedule.room_number})"
                    cell = ws.cell(row=row_idx, column=col, value=text)
                    cell.fill = PatternFill(start_color="FFD7D7", end_color="FFD7D7", fill_type="solid")
                else:
                    cell = ws.cell(row=row_idx, column=col, value="Free")
                    cell.fill = PatternFill(start_color="D7FFD7", end_color="D7FFD7", fill_type="solid")

                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = thin_border
                cell.font = Font(size=9)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=f"full_timetable_{datetime.now().strftime('%Y%m%d')}.xlsx",
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@app.route('/export/teacher/<int:id>')
@login_required
def export_teacher_timetable(id):
    """Export individual teacher timetable as Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    teacher = Teacher.query.get_or_404(id)
    wb = Workbook()
    ws = wb.active
    ws.title = teacher.name

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Title
    ws.merge_cells('A1:F1')
    title_cell = ws.cell(row=1, column=1, value=f"Timetable - {teacher.name}")
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal='center')

    if teacher.department:
        ws.merge_cells('A2:F2')
        dept_cell = ws.cell(row=2, column=1, value=f"Department: {teacher.department}")
        dept_cell.alignment = Alignment(horizontal='center')

    start_row = 4

    # Headers
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    ws.cell(row=start_row, column=1, value="Time Slot").font = header_font
    ws.cell(row=start_row, column=1).fill = header_fill
    ws.cell(row=start_row, column=1).border = thin_border
    ws.column_dimensions['A'].width = 16

    for col, day in enumerate(DAYS, start=2):
        cell = ws.cell(row=start_row, column=col, value=day)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = 20

    # Data
    for row_idx, ts in enumerate(TIME_SLOTS, start=start_row + 1):
        ws.cell(row=row_idx, column=1, value=ts).font = Font(bold=True, size=10)
        ws.cell(row=row_idx, column=1).border = thin_border

        for col, day in enumerate(DAYS, start=2):
            schedule = Schedule.query.filter_by(
                teacher_id=teacher.id, day=day, time_slot=ts
            ).first()

            if schedule and schedule.status == 'Engaged':
                text = schedule.subject or 'Engaged'
                if schedule.room_number:
                    text += f"\n({schedule.room_number})"
                cell = ws.cell(row=row_idx, column=col, value=text)
                cell.fill = PatternFill(start_color="FFD7D7", end_color="FFD7D7", fill_type="solid")
            else:
                cell = ws.cell(row=row_idx, column=col, value="Free")
                cell.fill = PatternFill(start_color="D7FFD7", end_color="D7FFD7", fill_type="solid")

            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border
            cell.font = Font(size=10)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    safe_name = teacher.name.replace(' ', '_')
    return send_file(
        output,
        as_attachment=True,
        download_name=f"timetable_{safe_name}_{datetime.now().strftime('%Y%m%d')}.xlsx",
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# ─── Admin Routes ────────────────────────────────────────────────────────────

@app.route('/users')
@admin_required
def manage_users():
    users = User.query.order_by(User.username).all()
    return render_template('users.html', users=users)

@app.route('/users/toggle_admin/<int:id>')
@admin_required
def toggle_admin(id):
    user = User.query.get_or_404(id)
    if user.is_admin and User.query.filter_by(is_admin=True).count() <= 1:
        flash('Cannot remove the last admin.', 'danger')
    else:
        user.is_admin = not user.is_admin
        db.session.commit()
        flash(f"User {user.username}'s admin status updated.", 'success')
    return redirect(url_for('manage_users'))

@app.route('/users/delete/<int:id>')
@admin_required
def delete_user(id):
    user = User.query.get_or_404(id)
    if user == current_user:
        flash('You cannot delete your own account.', 'danger')
    elif user.is_admin and User.query.filter_by(is_admin=True).count() <= 1:
        flash('Cannot delete the last admin.', 'danger')
    else:
        db.session.delete(user)
        db.session.commit()
        flash(f"User {user.username} has been deleted.", 'success')
    return redirect(url_for('manage_users'))

@app.route('/teacher/add', methods=['GET', 'POST'])
@admin_required
def add_teacher():
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        dept = request.form.get('department', '').strip() or None
        email = request.form.get('email', '').strip() or None
        phone = request.form.get('phone', '').strip() or None

        if not name:
            flash('Teacher name is required!', 'danger')
            return redirect(url_for('add_teacher'))

        if Teacher.query.filter_by(name=name).first():
            flash('Teacher already exists!', 'danger')
            return redirect(url_for('add_teacher'))

        teacher = Teacher(name=name, department=dept, email=email, phone=phone)
        db.session.add(teacher)

        for day in DAYS:
            for time_slot in TIME_SLOTS:
                schedule = Schedule(
                    day=day,
                    time_slot=time_slot,
                    status="Free",
                    room_number=None,
                    subject=None,
                    teacher=teacher
                )
                db.session.add(schedule)

        db.session.commit()
        flash('Teacher added successfully!', 'success')
        return redirect(url_for('index'))

    departments = db.session.query(Teacher.department).filter(
        Teacher.department.isnot(None)
    ).distinct().order_by(Teacher.department).all()
    departments = [d[0] for d in departments]
    return render_template('add_teacher.html', departments=departments)

@app.route('/teacher/edit/<int:id>', methods=['GET', 'POST'])
@admin_required
def edit_teacher(id):
    teacher = Teacher.query.get_or_404(id)

    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        dept = request.form.get('department', '').strip() or None
        email = request.form.get('email', '').strip() or None
        phone = request.form.get('phone', '').strip() or None

        existing = Teacher.query.filter_by(name=name).first()
        if existing and existing.id != id:
            flash('Teacher name already exists!', 'danger')
            return redirect(url_for('edit_teacher', id=id))

        teacher.name = name
        teacher.department = dept
        teacher.email = email
        teacher.phone = phone
        db.session.commit()
        flash('Teacher updated successfully!', 'success')
        return redirect(url_for('index'))

    departments = db.session.query(Teacher.department).filter(
        Teacher.department.isnot(None)
    ).distinct().order_by(Teacher.department).all()
    departments = [d[0] for d in departments]
    return render_template('edit_teacher.html', teacher=teacher, departments=departments)

@app.route('/teacher/delete/<int:id>')
@admin_required
def delete_teacher(id):
    teacher = Teacher.query.get_or_404(id)
    db.session.delete(teacher)
    db.session.commit()
    flash('Teacher deleted successfully!', 'success')
    return redirect(url_for('index'))

@app.route('/schedule/edit', methods=['GET', 'POST'])
@admin_required
def edit_schedule():
    if request.method == 'POST':
        teacher_id = request.form['teacher_id']
        day = request.form['day']
        time_slot = request.form['time_slot']
        status = request.form['status']
        room_number = request.form.get('room_number', '').strip() or None
        subject = request.form.get('subject', '').strip() or None

        schedule = Schedule.query.filter_by(
            teacher_id=teacher_id,
            day=day,
            time_slot=time_slot
        ).first()

        if schedule:
            # Room conflict detection
            if room_number and status == 'Engaged':
                conflict = Schedule.query.filter(
                    Schedule.room_number == room_number,
                    Schedule.day == day,
                    Schedule.time_slot == time_slot,
                    Schedule.status == 'Engaged',
                    Schedule.id != schedule.id
                ).first()
                if conflict:
                    flash(f'Warning: Room {room_number} is already booked by {conflict.teacher.name} at this time!', 'warning')

            schedule.status = status
            schedule.room_number = room_number
            schedule.subject = subject

            if status == 'Free':
                schedule.room_number = None
                schedule.subject = None

            db.session.commit()
            flash('Schedule updated successfully!', 'success')
        else:
            flash('Schedule not found!', 'danger')

        return redirect(url_for('timetable'))

    teacher_id = request.args.get('teacher_id')
    day = request.args.get('day')
    time_slot = request.args.get('time_slot')

    schedule = Schedule.query.filter_by(
        teacher_id=teacher_id,
        day=day,
        time_slot=time_slot
    ).first_or_404()

    return render_template('edit_schedule.html', schedule=schedule)

@app.route('/api/schedule/<int:teacher_id>/<string:day>/<string:time_slot>')
@csrf.exempt
@login_required
def get_schedule(teacher_id, day, time_slot):
    schedule = Schedule.query.filter_by(
        teacher_id=teacher_id,
        day=day,
        time_slot=time_slot
    ).first()

    if schedule:
        return jsonify({
            'status': schedule.status,
            'room_number': schedule.room_number,
            'subject': schedule.subject
        })
    else:
        return jsonify({
            'status': 'Free',
            'room_number': None,
            'subject': None
        })

@app.route('/api/room_conflict/<string:room>/<string:day>/<string:time_slot>')
@csrf.exempt
@login_required
def check_room_conflict(room, day, time_slot):
    """Check if a room is already booked for a given day and time slot."""
    conflicts = Schedule.query.filter_by(
        room_number=room,
        day=day,
        time_slot=time_slot,
        status='Engaged'
    ).all()

    return jsonify({
        'has_conflict': len(conflicts) > 0,
        'booked_by': [{'teacher': c.teacher.name, 'subject': c.subject} for c in conflicts]
    })

# ─── Bulk Import Routes ─────────────────────────────────────────────────────

@app.route('/import/teachers', methods=['GET', 'POST'])
@admin_required
def import_teachers():
    if request.method == 'POST':
        if 'file' not in request.files:
            flash('No file uploaded.', 'danger')
            return redirect(url_for('import_teachers'))

        file = request.files['file']
        if file.filename == '':
            flash('No file selected.', 'danger')
            return redirect(url_for('import_teachers'))

        if not file.filename.endswith(('.xlsx', '.xls', '.csv')):
            flash('Please upload an Excel (.xlsx/.xls) or CSV (.csv) file.', 'danger')
            return redirect(url_for('import_teachers'))

        try:
            import pandas as pd

            if file.filename.endswith('.csv'):
                df = pd.read_csv(file)
            else:
                df = pd.read_excel(file)

            # Normalize column names
            df.columns = [c.strip().lower() for c in df.columns]

            required_col = 'name'
            if required_col not in df.columns:
                flash('File must have a "name" column.', 'danger')
                return redirect(url_for('import_teachers'))

            added = 0
            skipped = 0
            for _, row in df.iterrows():
                name = str(row.get('name', '')).strip()
                if not name:
                    skipped += 1
                    continue

                if Teacher.query.filter_by(name=name).first():
                    skipped += 1
                    continue

                dept = row.get('department')
                email = row.get('email')
                phone = row.get('phone')

                # Use pandas-aware NaN check
                dept = None if pd.isna(dept) else str(dept).strip() or None
                email = None if pd.isna(email) else str(email).strip() or None
                phone = None if pd.isna(phone) else str(phone).strip() or None

                teacher = Teacher(name=name, department=dept, email=email, phone=phone)
                db.session.add(teacher)

                for day in DAYS:
                    for ts in TIME_SLOTS:
                        schedule = Schedule(
                            day=day, time_slot=ts, status="Free",
                            room_number=None, subject=None, teacher=teacher
                        )
                        db.session.add(schedule)
                added += 1

            db.session.commit()
            flash(f'Import complete! Added {added} teachers, skipped {skipped}.', 'success')
            return redirect(url_for('index'))

        except Exception as e:
            flash(f'Error processing file: {str(e)}', 'danger')
            return redirect(url_for('import_teachers'))

    return render_template('import_teachers.html')

# ─── Administration Routes ───────────────────────────────────────────────────

@app.route('/initialize-admin', methods=['GET'])
@csrf.exempt
def initialize_admin():
    create_admin_user()
    flash('Admin user created or reset with username: admin and password: admin123', 'success')
    return redirect(url_for('login'))

@app.route('/reset-database', methods=['GET'])
@csrf.exempt
def reset_db_route():
    try:
        db.drop_all()
        db.create_all()
        init_data()
        flash('Database has been reset. Admin user created with username: admin and password: admin123', 'success')
        return redirect(url_for('login'))
    except Exception as e:
        return f"Error resetting database: {str(e)}", 500

# ─── Main Entry Point ───────────────────────────────────────────────────────

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
        init_data()
    app.run(debug=os.environ.get('FLASK_DEBUG', 'false').lower() == 'true', host='0.0.0.0', port=5000)
